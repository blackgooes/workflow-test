"""读取 inbox/ 下的一个文件（txt、md、pdf、docx），在末尾加一段 hello，生成 outbox/<同名>.docx。

用法：python make_hello.py inbox/<文件名>
"""
import sys
from pathlib import Path

from docx import Document

ALLOWED = {".txt", ".md", ".pdf", ".docx"}


def read_text(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace").splitlines()
    if ext == ".docx":
        return [p.text for p in Document(str(path)).paragraphs]
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages).splitlines()
    raise ValueError(f"不支持的文件类型：{path.name}")


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("用法：python make_hello.py inbox/<文件名>")
    src = Path(sys.argv[1])
    if src.parts[0] != "inbox" or ".." in src.parts or src.suffix.lower() not in ALLOWED:
        sys.exit(f"只处理 inbox/ 下的 txt、md、pdf、docx：{src}")

    lines = [line for line in read_text(src) if line.strip()]
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.add_paragraph("hello")

    out = Path("outbox") / f"{src.stem}.docx"
    out.parent.mkdir(exist_ok=True)
    doc.save(out)
    print(f"{src} -> {out}（原文 {len(lines)} 段，末尾已加 hello）")


if __name__ == "__main__":
    main()
